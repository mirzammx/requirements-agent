"""Local web app: buttons for shortlist -> fetch -> triage -> decision -> response -> download.

Run in Docker:  docker compose up web   then open http://localhost:8010

Who decides what:
- BD, on a shortlist: Fetch documents / Skip for each candidate tender (no bid decision yet).
- Mahadevan, on a bid's triage card: Go / No-go / Ask, after the full RFQ has been read.
"""

import io
import json
import os
import re
import shutil
import threading
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

from flask import (Flask, abort, jsonify, redirect, render_template, request, send_file, send_from_directory,
                   url_for)
from werkzeug.utils import secure_filename

import draft
import shortlist
import triage
from bidtriage.email_loader import load_bid
from bidtriage.intake import classify_mail, load_intake_settings
from bidtriage.render import CSS
from bidtriage.reply import build_reply
from bidtriage.schema import BidAssessment
from bidtriage.shortlist_out import EXTRA_CSS
from bidtriage.web import jobs

ROOT = Path(__file__).resolve().parent
ID_RE = re.compile(r"^\d{8}-[0-9a-f]{6}$")
TENDER_ACTIONS = ("Fetch", "Skip")          # BD, on the shortlist
BID_DECISIONS = ("Go", "No-go", "Ask")       # Mahadevan, on the triage card
CANDIDATE_CATEGORIES = ("sbu2_power", "manual_check")
EMAIL_SUFFIXES = {".msg", ".eml"}
STEP_TEXT = {
    "triage": "Reading the documents and writing the triage card (usually 1–2 minutes)…",
    "draft": "Drafting the response pack (about 5–10 minutes on Gemini)…",
    "shortlist": "Screening the tenders (a few minutes for 100 tenders)…",
}
FILE_NOTES = {"Technical_Compliance.docx": "Mahadevan: resolve every yellow REVIEW / DEVIATION",
              "Compliance_Statement.docx": "Buyer's compliance form",
              "Price_Bid.xlsx": "BD: enter unit prices and GST % in the yellow cells",
              "Declarations.docx": "Print on letterhead and sign",
              "PreBid_Queries.docx": "Clarification letter to the buyer",
              "Checklist.md": "Every document the RFQ requires"}


def _unique_name(original: str, taken: set[str]) -> str:
    """A safe file name that keeps the extension (even for non-English names) and never overwrites."""
    suffix = Path(original).suffix.lower()
    suffix = suffix if suffix[1:].isalnum() else ""
    stem = secure_filename(Path(original).stem) or "upload"
    name, n = f"{stem}{suffix}", 2
    while name in taken:
        name, n = f"{stem}_{n}{suffix}", n + 1
    return name


def _new_id() -> str:
    return f"{datetime.now():%Y%m%d}-{uuid.uuid4().hex[:6]}"


def _read_json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


def _default_provider() -> str:
    return "gemini" if os.environ.get("GEMINI_API_KEY") and not os.environ.get("ANTHROPIC_API_KEY") else "claude"


def _days_key(item: dict):
    d = item.get("days_left")
    return (0, d) if d is not None and d >= 0 else (1, -(d or 0)) if d is not None else (2, 0)


def create_app(data_dir: Path = ROOT / "data", samples_dir: Path = ROOT / "samples", sync_jobs: bool = False,
               allowed_hosts: set[str] = frozenset({"localhost:8010", "127.0.0.1:8010"}),
               intake_settings: dict | None = None):
    app = Flask(__name__, template_folder=str(ROOT / "bidtriage" / "web" / "templates"))
    app.config.update(MAX_CONTENT_LENGTH=50 * 1024 * 1024, SYNC_JOBS=sync_jobs)
    data_dir = Path(data_dir)
    bids_dir, lists_dir = data_dir / "bids", data_dir / "shortlists"
    settings = intake_settings or load_intake_settings()
    for d in (bids_dir, lists_dir):
        if d.is_dir():
            jobs.recover([p for p in d.iterdir() if p.is_dir()])
    allowed_origins = {f"http://{h}" for h in allowed_hosts}

    @app.before_request
    def same_origin_only():
        # Another website open in the browser must not drive this app (paid AI calls) or read it
        # (DNS rebinding): require our own Host, and our own Origin/Referer on any POST that sends one.
        if request.host not in allowed_hosts:
            abort(400)
        if request.method == "POST":
            origin = request.headers.get("Origin") or ""
            referer = request.headers.get("Referer") or ""
            if origin and origin not in allowed_origins:
                abort(403)
            if not origin and referer and not any(referer.startswith(o + "/") for o in allowed_origins):
                abort(403)

    @app.context_processor
    def common():
        return {"css": CSS + EXTRA_CSS, "providers": ("claude", "gemini"), "default_provider": _default_provider()}

    # ---------------------------------------------------------------- data helpers

    def folder(kind_dir: Path, item_id: str) -> Path:
        path = kind_dir / item_id
        if not ID_RE.match(item_id) or not path.is_dir():
            abort(404)
        return path

    def samples() -> list[str]:
        if not Path(samples_dir).is_dir():
            return []
        return sorted(p.name for p in Path(samples_dir).iterdir() if p.is_dir())

    def shortlist_items(run: Path) -> list[dict]:
        return _read_json(run / "result" / "shortlist.json", [])

    def bid_state(bid: Path) -> dict:
        """Everything the pages need to know about where a bid stands."""
        meta = _read_json(bid / "meta.json", {"id": bid.name, "title": bid.name})
        response = bid / "card" / "response"
        summary = _read_json(response / "summary.json", {})
        files = [{"name": n, "count": c, "note": FILE_NOTES.get(n, ""), "exists": (response / n).is_file()}
                 for n, c in summary.items() if c is not None or n == "Checklist.md"]
        state = {
            "meta": meta,
            "status": jobs.read_status(bid),
            "has_card": (bid / "card" / "card.html").is_file(),
            "can_draft": (bid / "card" / "assessment.json").is_file(),
            "decision": _read_json(bid / "decision.json", {}).get("decision", ""),
            "files": files,
            "has_files": any(f["exists"] for f in files),
            "has_saved_draft": (response / "draft.json").is_file(),
            "mail": meta.get("mail"),
        }
        state["pack_current"] = state["has_files"] and _pack_is_current(bid)
        state["reply"] = _reply_state(bid, state)
        state["steps"], state["next_action"] = _bid_steps(state)
        return state

    def listing(kind_dir: Path) -> list[dict]:
        if not kind_dir.is_dir():
            return []
        items = [{**_read_json(p / "meta.json", {"id": p.name, "title": p.name}), "status": jobs.read_status(p)}
                 for p in kind_dir.iterdir() if p.is_dir() and ID_RE.match(p.name)]
        return sorted(items, key=lambda m: m.get("created", ""), reverse=True)

    def all_bids() -> list[dict]:
        if not bids_dir.is_dir():
            return []
        states = [bid_state(p) for p in bids_dir.iterdir() if p.is_dir() and ID_RE.match(p.name)]
        return sorted(states, key=lambda s: s["meta"].get("created", ""), reverse=True)

    def bids_for_tenders(run_id: str) -> dict[str, str]:
        """tender_id -> bid id, for bids started from this shortlist."""
        found = {}
        for s in all_bids():
            source = s["meta"].get("source") or {}
            if source.get("shortlist") == run_id and source.get("tender_id"):
                found.setdefault(source["tender_id"], s["meta"]["id"])
        return found

    def needs_action() -> list[dict]:
        actions = []
        for run in listing(lists_dir):
            run_dir = lists_dir / run["id"]
            if run["status"].get("state") != "done":
                continue
            decisions = _read_json(run_dir / "decisions.json", {})
            started = bids_for_tenders(run["id"])
            for t in sorted((i for i in shortlist_items(run_dir) if i["category"] in CANDIDATE_CATEGORIES),
                            key=_days_key):
                if t.get("days_left") is not None and t["days_left"] < 0:
                    continue  # closed: nothing to do
                choice = decisions.get(t["tender_id"])
                if not choice:
                    actions.append({"who": "BD", "what": "Fetch documents or skip", "title": t["summary"],
                                    "detail": f"{t['authority']} · {_closes(t)}",
                                    "href": url_for("show_shortlist", run_id=run["id"]) + "#" + _anchor(t["tender_id"]),
                                    "button": "Open shortlist"})
                elif choice == "Fetch" and t["tender_id"] not in started:
                    actions.append({"who": "BD", "what": "Start the bid", "title": t["summary"],
                                    "detail": f"{t['authority']} · {_closes(t)} · upload the RFQ documents",
                                    "href": url_for("new_bid", shortlist=run["id"], tender=t["tender_id"]),
                                    "button": "Start bid"})
        for s in all_bids():
            m, st = s["meta"], s["status"].get("state")
            href = url_for("show_bid", bid_id=m["id"])
            if st == "running":
                actions.append({"who": "Wait", "what": STEP_TEXT.get(s["status"].get("step", ""), "Working…"),
                                "title": m["title"], "detail": "", "href": href, "button": "Open bid"})
            elif st == "failed":
                actions.append({"who": "Retry", "what": f"The {s['status'].get('step')} step failed",
                                "title": m["title"], "detail": s["status"].get("error", ""), "href": href,
                                "button": "Open bid"})
            elif s["can_draft"] and not s["decision"]:
                actions.append({"who": "Mahadevan", "what": "Go / No-go on the card", "title": m["title"],
                                "detail": "Triage card ready", "href": href + "#decide", "button": "Read card"})
            elif s["decision"] == "Go" and not s["has_files"]:
                actions.append({"who": "BD", "what": "Generate the response", "title": m["title"],
                                "detail": "Decision: Go", "href": href + "#response", "button": "Open bid"})
            elif s["reply"]["state"] == "due":
                actions.append({"who": "You", "what": "Reply to BD is ready to draft", "title": m["title"],
                                "detail": "Start the mail watcher (mailwatch.ps1) so it lands in your Outlook Drafts",
                                "href": href + "#decide", "button": "Open bid"})
            elif s["reply"]["state"] == "drafted" and _within_days(s["reply"].get("drafted_at", ""), 3):
                actions.append({"who": "You", "what": "Check and send the reply from Outlook Drafts",
                                "title": m["title"], "detail": f"Drafted {s['reply']['drafted_at']} "
                                f"({s['reply']['decision']})", "href": href + "#decide", "button": "Open bid"})
        return actions

    def provider_arg() -> str:
        p = request.form.get("provider", "")
        return p if p in ("claude", "gemini") else _default_provider()

    def tender_from(run_id: str, tender_id: str) -> dict | None:
        if not run_id or not tender_id or not ID_RE.match(run_id) or not (lists_dir / run_id).is_dir():
            return None
        return next((t for t in shortlist_items(lists_dir / run_id) if t["tender_id"] == tender_id), None)

    # ---------------------------------------------------------------- overview + lists

    @app.get("/")
    def index():
        return render_template("index.html", active="overview", actions=needs_action(),
                               shortlists=listing(lists_dir)[:5], bids=all_bids()[:5])

    @app.get("/shortlists")
    def list_shortlists(error: str = "", status: int = 200):
        return render_template("shortlists.html", active="shortlists", shortlists=listing(lists_dir),
                               error=error), status

    @app.get("/bids")
    def list_bids():
        return render_template("bids.html", active="bids", bids=all_bids())

    @app.get("/bids/new")
    def new_bid(error: str = "", status: int = 200, form: dict | None = None):
        form = form or {}
        run_id = request.values.get("shortlist", "")
        tender_id = request.values.get("tender", "")
        tender = tender_from(run_id, tender_id)
        if tender and not form:
            form = {"subject": tender["summary"],
                    "body": (f"Tender: {tender.get('tender_no') or tender['tender_id']}\n"
                             f"Authority: {tender['authority']}\nLocation: {tender.get('location', '')}\n"
                             f"Closes: {tender.get('closing_date', '')}\nValue: {tender.get('amount') or '-'}\n"
                             f"From the BidAssist shortlist; BD fetched the documents.")}
        return render_template("bid_new.html", active="new", samples=samples(), error=error, form=form,
                               tender=tender, run_id=run_id if tender else "",
                               run_meta=_read_json(lists_dir / run_id / "meta.json", {}) if tender else {}), status

    # ---------------------------------------------------------------- bids

    @app.post("/bids")
    def create_bid():
        provider = provider_arg()
        sample = request.form.get("sample", "")
        files = [f for f in request.files.getlist("files") if f and f.filename]
        subject, body = request.form.get("subject", "").strip(), request.form.get("body", "").strip()
        run_id, tender_id = request.form.get("shortlist", ""), request.form.get("tender", "")
        form = {"subject": subject, "body": body}
        if sample and sample not in samples():
            abort(400)
        if (run_id or tender_id) and not tender_from(run_id, tender_id):
            abort(400)
        if not sample and not files:
            return new_bid("Add the bid documents (email .msg/.eml, PDFs, Excel) or pick a sample folder.", 400, form)
        if sum(Path(f.filename).suffix.lower() in EMAIL_SUFFIXES for f in files) > 1:
            return new_bid("Upload one email file (.msg or .eml) per bid; add its documents alongside it.", 400, form)

        bid_id = _new_id()
        bid = bids_dir / bid_id
        inp = bid / "input"
        inp.mkdir(parents=True)
        if sample:
            shutil.copytree(Path(samples_dir) / sample, inp, dirs_exist_ok=True)
        has_email = any(Path(f.filename).suffix.lower() in EMAIL_SUFFIXES for f in files)
        typed = (subject or body) and not has_email
        taken = {p.name for p in inp.iterdir()} | ({"email.txt"} if typed else set())
        for f in files:
            name = _unique_name(f.filename, taken)
            taken.add(name)
            f.save(inp / name)
        if typed and not (inp / "email.txt").exists():
            (inp / "email.txt").write_text(f"Subject: {subject}\n\n{body}\n", encoding="utf-8")

        title = subject or sample or (files[0].filename if files else bid_id)
        extra = {}
        if run_id:
            extra["source"] = {"shortlist": run_id, "tender_id": tender_id,
                               "shortlist_created": _read_json(lists_dir / run_id / "meta.json", {}).get("created", "")}
        start_bid(bid, provider, title, extra)
        return redirect(url_for("show_bid", bid_id=bid_id))

    def start_bid(bid: Path, provider: str, title: str, extra: dict) -> None:
        """Write meta.json and start triage on bid/input (one email file plus documents, or documents)."""
        bid_input = bid / "input"
        meta = {"id": bid.name, "title": title, "created": f"{datetime.now():%Y-%m-%d %H:%M}", "provider": provider,
                "input": "input", **extra}
        (bid / "meta.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
        jobs.start(bid, "triage", lambda: triage.run_triage(bid_input, provider=provider, out_dir=bid / "card"),
                   sync=app.config["SYNC_JOBS"])

    def start_shortlist(run: Path, paths: list[Path], provider: str, title: str, extra: dict) -> None:
        meta = {"id": run.name, "title": title, "created": f"{datetime.now():%Y-%m-%d %H:%M}", "provider": provider,
                **extra}
        (run / "meta.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
        jobs.start(run, "shortlist",
                   lambda: shortlist.run_shortlist(paths, provider=provider, out_dir=run / "result"),
                   sync=app.config["SYNC_JOBS"])

    # ---------------------------------------------------------------- mail watcher API

    @app.post("/api/intake")
    def api_intake():
        """The watcher posts each new mail from the Outlook "Bid desk" folder here, once per EntryID."""
        upload = request.files.get("file")
        entry_id = request.form.get("entry_id", "").strip()
        if not upload or not entry_id or Path(upload.filename or "").suffix.lower() not in EMAIL_SUFFIXES:
            abort(400)
        with INTAKE_LOCK:
            log = _read_json(data_dir / "intake.json", {})
            if entry_id in log:
                return jsonify(log[entry_id])
            provider = _default_provider()
            item_id = _new_id()
            staging = data_dir / "intake" / item_id
            staging.mkdir(parents=True)
            mail_path = staging / _unique_name(upload.filename, set())
            upload.save(mail_path)
            try:
                mail = load_bid(mail_path)
            except Exception as exc:
                shutil.rmtree(staging, ignore_errors=True)
                return jsonify({"error": f"could not read the mail: {exc}"}), 400
            mail_meta = {"entry_id": entry_id, "subject": mail.subject, "sender": mail.sender,
                         "received": request.form.get("received", "") or mail.date}
            kind, screen, ignored = classify_mail(mail, settings)
            if kind == "bid":
                bid = bids_dir / item_id
                (bid / "input").mkdir(parents=True)
                shutil.move(str(mail_path), bid / "input" / mail_path.name)
                shutil.rmtree(staging, ignore_errors=True)
                result = {"kind": "bid", "bid_id": item_id}
                log[entry_id] = result
                _atomic_write_text(data_dir / "intake.json", json.dumps(log, indent=2))
                start_bid(bid, provider, mail.subject or mail_path.stem, {"mail": mail_meta})
            elif screen:
                run = lists_dir / item_id
                (run / "input").mkdir(parents=True)
                paths, taken = [], set()
                for a in screen:
                    name = _unique_name(a.filename, taken)
                    taken.add(name)
                    (run / "input" / name).write_bytes(a.data)
                    paths.append(run / "input" / name)
                shutil.rmtree(staging, ignore_errors=True)
                result = {"kind": "bidassist", "shortlist_id": item_id, "ignored": ignored}
                log[entry_id] = result
                _atomic_write_text(data_dir / "intake.json", json.dumps(log, indent=2))
                start_shortlist(run, paths, provider, ", ".join(a.filename for a in screen), {"mail": mail_meta})
            else:
                shutil.rmtree(staging, ignore_errors=True)
                result = {"kind": "ignored", "ignored": ignored,
                          "reason": f"no BidAssist alert listed in screen_alerts {settings['screen_alerts']}"}
                log[entry_id] = result
                _atomic_write_text(data_dir / "intake.json", json.dumps(log, indent=2))
        return jsonify(result)

    @app.get("/api/replies")
    def api_replies():
        """Replies the watcher should draft in Outlook now (never sent automatically)."""
        due = []
        for s in all_bids():
            if s["reply"]["state"] != "due":
                continue
            bid = bids_dir / s["meta"]["id"]
            try:
                assessment = BidAssessment.model_validate_json(
                    (bid / "card" / "assessment.json").read_text(encoding="utf-8"))
            except (OSError, ValueError):
                continue
            attachments, placeholders = [], None
            if s["decision"] == "Go":
                placeholders = sum(f["count"] or 0 for f in s["files"] if f["name"].endswith(".docx"))
                zip_path = _write_pack_zip(bid, s["meta"].get("title", bid.name), response_files(bid))
                attachments.append(zip_path.relative_to(data_dir).as_posix())
            reply = build_reply(assessment, s["decision"], pack_placeholders=placeholders,
                                signature=settings["signature"])
            due.append({"bid_id": bid.name, "entry_id": s["mail"]["entry_id"], "decision": s["decision"],
                        "html": reply["html"], "text": reply["text"], "cc": settings["mahadevan_email"],
                        "attachments": attachments})
        return jsonify(due)

    @app.post("/api/replies/<bid_id>/drafted")
    def api_reply_drafted(bid_id):
        bid = folder(bids_dir, bid_id)
        drafted = request.form.get("decision", "")
        with DECISION_LOCK:
            current = _read_json(bid / "decision.json", {}).get("decision", "")
            if not current or drafted not in BID_DECISIONS:
                abort(400)
            if drafted != current:
                # Mahadevan changed his mind while the draft was being made: the saved draft is outdated.
                return jsonify({"ok": False, "current": current,
                                "error": f"decision is now {current}; delete the {drafted} draft"}), 409
            _atomic_write_text(bid / "reply.json", json.dumps(
                {"decision": drafted, "drafted_at": f"{datetime.now():%Y-%m-%d %H:%M}"}, indent=2))
        return jsonify({"ok": True})

    @app.get("/bids/<bid_id>")
    def show_bid(bid_id):
        bid = folder(bids_dir, bid_id)
        s = bid_state(bid)
        return render_template("bid.html", active="bids", **s,
                               step_text=STEP_TEXT.get(s["status"].get("step", ""), "Working…"),
                               msg=request.args.get("msg", ""))

    @app.get("/bids/<bid_id>/card")
    def bid_card(bid_id):
        card = folder(bids_dir, bid_id) / "card" / "card.html"
        if not card.is_file():
            abort(404)
        return send_file(card)

    @app.post("/bids/<bid_id>/decide")
    def decide_bid(bid_id):
        bid = folder(bids_dir, bid_id)
        decision = request.form.get("decision", "")
        if decision not in BID_DECISIONS or not (bid / "card" / "assessment.json").is_file():
            abort(400)
        with DECISION_LOCK:
            _atomic_write_text(bid / "decision.json", json.dumps(
                {"decision": decision, "at": f"{datetime.now():%Y-%m-%d %H:%M}"}, indent=2))
        if decision == "Go" and not (response_files(bid) and _pack_is_current(bid)):
            # Go means BD needs the response pack (a fresh one if the card changed): start it now;
            # the reply waits for it.
            meta = _read_json(bid / "meta.json", {})
            bid_input = bid / meta.get("input", "input")
            provider = meta.get("provider") or _default_provider()
            jobs.start(bid, "draft", lambda: draft.run_draft(bid_input, bid / "card", provider=provider),
                       sync=app.config["SYNC_JOBS"])
        return redirect(url_for("show_bid", bid_id=bid_id) + ("#response" if decision == "Go" else "#decide"))

    @app.post("/bids/<bid_id>/triage")
    def retriage_bid(bid_id):
        bid = folder(bids_dir, bid_id)
        bid_input = bid / _read_json(bid / "meta.json", {}).get("input", "input")
        provider = provider_arg()
        started = jobs.start(bid, "triage",
                             lambda: triage.run_triage(bid_input, provider=provider, out_dir=bid / "card"),
                             sync=app.config["SYNC_JOBS"])
        return redirect(url_for("show_bid", bid_id=bid_id, msg="" if started else "A job is already running."))

    @app.post("/bids/<bid_id>/draft")
    def draft_bid(bid_id):
        bid = folder(bids_dir, bid_id)
        if not (bid / "card" / "assessment.json").is_file():
            abort(400)
        meta = _read_json(bid / "meta.json", {})
        bid_input = bid / meta.get("input", "input")
        provider = provider_arg()
        started = jobs.start(bid, "draft",
                             lambda: draft.run_draft(bid_input, bid / "card", provider=provider),
                             sync=app.config["SYNC_JOBS"])
        return redirect(url_for("show_bid", bid_id=bid_id, msg="" if started else "A job is already running."))

    @app.post("/bids/<bid_id>/rebuild")
    def rebuild_bid(bid_id):
        bid = folder(bids_dir, bid_id)
        saved = bid / "card" / "response" / "draft.json"
        if not saved.is_file():
            abort(400)
        bid_input = bid / _read_json(bid / "meta.json", {}).get("input", "input")
        started = jobs.start(bid, "draft",
                             lambda: draft.run_draft(bid_input, bid / "card", from_draft=saved),
                             sync=app.config["SYNC_JOBS"])
        return redirect(url_for("show_bid", bid_id=bid_id, msg="" if started else "A job is already running."))

    def response_files(bid: Path) -> list[str]:
        """Files of the current pack only: a document the latest draft did not need is never offered."""
        summary = _read_json(bid / "card" / "response" / "summary.json", {})
        return [n for n, c in summary.items()
                if (c is not None or n == "Checklist.md") and (bid / "card" / "response" / n).is_file()]

    @app.get("/bids/<bid_id>/files/<name>")
    def bid_file(bid_id, name):
        bid = folder(bids_dir, bid_id)
        if name not in response_files(bid):
            abort(404)
        return send_from_directory(bid / "card" / "response", name, as_attachment=True)

    @app.get("/bids/<bid_id>/response.zip")
    def bid_zip(bid_id):
        bid = folder(bids_dir, bid_id)
        names = response_files(bid)
        if not names:
            abort(404)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for n in names:
                z.write(bid / "card" / "response" / n, n)
        buf.seek(0)
        title = secure_filename(_read_json(bid / "meta.json", {}).get("title", bid_id)) or bid_id
        return send_file(buf, mimetype="application/zip", as_attachment=True, download_name=f"{title}_response.zip")

    # ---------------------------------------------------------------- shortlists

    @app.post("/shortlists")
    def create_shortlist():
        provider = provider_arg()
        files = [f for f in request.files.getlist("files") if f and f.filename]
        if not files:
            return list_shortlists("Choose one or more BidAssist 'Tender List' .xlsx files.", 400)
        run_id = _new_id()
        run = lists_dir / run_id
        inp = run / "input"
        inp.mkdir(parents=True)
        paths, taken = [], set()
        for f in files:
            name = _unique_name(f.filename, taken)
            taken.add(name)
            f.save(inp / name)
            paths.append(inp / name)
        meta = {"id": run_id, "title": ", ".join(f.filename for f in files),
                "created": f"{datetime.now():%Y-%m-%d %H:%M}", "provider": provider}
        (run / "meta.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
        jobs.start(run, "shortlist",
                   lambda: shortlist.run_shortlist(paths, provider=provider, out_dir=run / "result"),
                   sync=app.config["SYNC_JOBS"])
        return redirect(url_for("show_shortlist", run_id=run_id))

    @app.get("/shortlists/<run_id>")
    def show_shortlist(run_id):
        run = folder(lists_dir, run_id)
        items = shortlist_items(run)
        groups = {c: sorted([i for i in items if i["category"] == c], key=_days_key)
                  for c in ("sbu2_power", "manual_check", "other_product_line", "not_relevant")}
        return render_template("shortlist.html", active="shortlists",
                               meta=_read_json(run / "meta.json", {"id": run_id}), status=jobs.read_status(run),
                               step_text=STEP_TEXT["shortlist"], groups=groups,
                               decisions=_read_json(run / "decisions.json", {}), bids=bids_for_tenders(run_id),
                               total=len(items), msg=request.args.get("msg", ""))

    @app.post("/shortlists/<run_id>/decide")
    def decide(run_id):
        run = folder(lists_dir, run_id)
        tender_id, decision = request.form.get("tender_id", ""), request.form.get("decision", "")
        if decision not in TENDER_ACTIONS or tender_id not in {i["tender_id"] for i in shortlist_items(run)}:
            abort(400)
        msg = ""
        with DECISION_LOCK:  # two quick clicks must not lose a decision or interleave workbook saves
            decisions = _read_json(run / "decisions.json", {})
            decisions[tender_id] = decision
            _atomic_write_text(run / "decisions.json", json.dumps(decisions, indent=2))
            try:
                _write_decision_to_xlsx(run / "result" / "shortlist.xlsx", tender_id, decision)
            except Exception:
                msg = ("Saved, but shortlist.xlsx could not be updated (is it open in Excel?). "
                       "Close it and click again to copy it into the workbook.")
        return redirect(url_for("show_shortlist", run_id=run_id, msg=msg) + "#" + _anchor(tender_id))

    @app.get("/shortlists/<run_id>/files/<name>")
    def shortlist_file(run_id, name):
        run = folder(lists_dir, run_id)
        if name not in ("digest.html", "shortlist.xlsx") or not (run / "result" / name).is_file():
            abort(404)
        return send_from_directory(run / "result", name, as_attachment=name.endswith(".xlsx"))

    app.add_template_filter(_anchor, "anchor")
    app.add_template_filter(_closes, "closes")
    return app


def _bid_steps(s: dict) -> tuple[list[dict], dict]:
    """The bid's five steps, each done / current / todo, and the one next action to show."""
    running, failed = s["status"].get("state") == "running", s["status"].get("state") == "failed"
    step = s["status"].get("step")
    raw = [
        ("Documents", True),
        ("Triage card", s["can_draft"]),
        ("Decide on the card", bool(s["decision"])),
        ("Response pack", s["has_files"]),
        ("Download the pack", False),
    ]
    current = next((i for i, (_, done) in enumerate(raw) if not done), len(raw) - 1)
    if s["has_files"]:
        current = 4
    steps = [{"label": label, "state": "done" if i < current else "current" if i == current else "todo"}
             for i, (label, _) in enumerate(raw)]
    if running:
        nxt = {"text": STEP_TEXT.get(step, "Working…"), "href": "", "button": ""}
    elif failed:
        nxt = {"text": f"The {step} step failed. Try it again below.", "href": "#status", "button": "See the error"}
    elif current == 2:
        nxt = {"text": "Mahadevan: read the triage card and decide Go / No-go / Ask.", "href": "#decide",
               "button": "Decide"}
    elif current == 3:
        nxt = {"text": "Generate the response pack (usually after a Go).", "href": "#response",
               "button": "Generate response"}
    elif current == 4:
        nxt = {"text": "Download the pack, complete the yellow placeholders, price and submit.",
               "href": "#response", "button": "Go to downloads"}
    else:
        nxt = {"text": "Waiting for the triage card.", "href": "", "button": ""}
    return steps, nxt


def _reply_state(bid: Path, s: dict) -> dict:
    """n/a (bid not from mail) · none (no decision) · waiting (Go, pack not ready) · due · drafted."""
    if not s.get("mail"):
        return {"state": "n/a"}
    if not s["decision"]:
        return {"state": "none"}
    drafted = _read_json(bid / "reply.json", {})
    if drafted.get("decision") == s["decision"]:
        return {"state": "drafted", **drafted}
    if s["decision"] == "Go" and (not s["pack_current"] or s["status"].get("state") == "running"):
        return {"state": "waiting"}
    return {"state": "due"}


def _pack_is_current(bid: Path) -> bool:
    """A response pack counts only if it was made from the current card (not before a re-triage)."""
    summary, card = bid / "card" / "response" / "summary.json", bid / "card" / "assessment.json"
    try:
        return summary.stat().st_mtime >= card.stat().st_mtime
    except OSError:
        return False


def _write_pack_zip(bid: Path, title: str, names: list[str]) -> Path:
    response = bid / "card" / "response"
    path = response / f"{secure_filename(title) or bid.name}_response.zip"
    tmp = path.with_suffix(".zip.tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.write(response / n, n)
    os.replace(tmp, path)
    return path


def _within_days(stamp: str, days: int) -> bool:
    try:
        return (datetime.now() - datetime.strptime(stamp, "%Y-%m-%d %H:%M")).days < days
    except ValueError:
        return False


INTAKE_LOCK = threading.Lock()


def _anchor(tender_id: str) -> str:
    return "t-" + re.sub(r"[^A-Za-z0-9]+", "-", tender_id).strip("-")


def _closes(t: dict) -> str:
    d = t.get("days_left")
    if d is None:
        return f"closes {t.get('closing_date') or 'date not given'}"
    if d < 0:
        return f"closed {t.get('closing_date')}"
    return "closes today" if d == 0 else f"closes {t.get('closing_date')} ({d} days left)"


DECISION_LOCK = threading.Lock()


def _atomic_write_text(path: Path, text: str) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, path)


def _write_decision_to_xlsx(path: Path, tender_id: str, decision: str) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    id_col, decision_col = headers.index("Tender Id"), headers.index("Decision")
    for row in ws.iter_rows(min_row=2):
        if row[id_col].value == tender_id:
            row[decision_col].value = decision
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)


if __name__ == "__main__":
    create_app().run(host="0.0.0.0", port=8000, debug=False, threaded=True)
