"""Digest page and tracker workbook for a classified BidAssist shortlist."""

from collections import defaultdict
from datetime import date
from html import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from bidtriage.render import CSS
from bidtriage.shortlister import Shortlisted

CATEGORY_ORDER = {"sbu2_power": 0, "manual_check": 1, "other_product_line": 2, "not_relevant": 3}
EXTRA_CSS = """
.tender{border-left:4px solid var(--ok);padding:10px 14px;margin:10px 0;border-radius:6px;background:var(--bg)}
.tender.closed{border-color:var(--line);opacity:.7}.tender .title{font-weight:600}
.meta{font-size:13px;color:var(--muted)}.urgent{color:var(--high);font-weight:700}
.tag{display:inline-block;font-size:11px;font-weight:600;border-radius:4px;padding:1px 7px;margin-left:6px;
background:var(--grey-bg);color:var(--muted)}
.summary-row{display:flex;gap:12px;flex-wrap:wrap}.stat{background:var(--bg);border-radius:8px;padding:8px 14px}
.stat b{font-size:20px;display:block}
"""


def _sort_key(s: Shortlisted):
    d = s.days_left
    return (0, d) if d is not None and d >= 0 else (1, -(d or 0)) if d is not None else (2, 0)


def _deadline(s: Shortlisted) -> str:
    if s.days_left is None:
        return f"Closes {escape(s.tender.closing_date or 'date not given')}"
    if s.days_left < 0:
        return f"Closed {escape(s.tender.closing_date)}"
    cls = " class='urgent'" if s.days_left <= 3 else ""
    return f"<span{cls}>Closes {escape(s.tender.closing_date)} ({s.days_left} days left)</span>"


def _tender_card(s: Shortlisted) -> str:
    t, e = s.tender, escape
    tags = "".join(f"<span class='tag'>{x}</span>" for x in
                   (["corrigendum"] if t.corrigendum else []) + ([f"{e(s.confidence)} confidence"]
                                                                  if s.confidence else []))
    closed = " closed" if s.days_left is not None and s.days_left < 0 else ""
    return (f"<div class='tender{closed}'><div class='title'>{e(t.summary)}{tags}</div>"
            f"<div class='meta'>{e(t.authority)} · {e(t.location)} · {e(t.tender_no or t.tender_id)}</div>"
            f"<div class='meta'>{_deadline(s)} · Value {e(t.amount or '-')} · EMD {e(t.emd or '-')}</div>"
            f"<p style='margin:6px 0 0'>{e(s.reason)}</p></div>")


def _table(rows: list[Shortlisted], with_route: bool) -> str:
    head = ("<th>Team</th>" if with_route else "") + "<th>Tender</th><th>Authority</th><th>Closes</th><th>Reason</th>"
    body = "".join(
        "<tr>" + (f"<td>{escape(s.route_to)}</td>" if with_route else "")
        + f"<td>{escape(s.tender.summary)}<br><span class='src'>{escape(s.tender.tender_no or s.tender.tender_id)}"
          f"</span></td><td>{escape(s.tender.authority)}</td><td>{escape(s.tender.closing_date)}</td>"
          f"<td>{escape(s.reason)}</td></tr>"
        for s in rows)
    return f"<div class='tablewrap'><table><tr>{head}</tr>{body}</table></div>"


def render_digest(items: list[Shortlisted], today: date) -> str:
    groups: dict[str, list[Shortlisted]] = defaultdict(list)
    for s in sorted(items, key=_sort_key):
        groups[s.category].append(s)
    power, manual = groups["sbu2_power"], groups["manual_check"]
    other, irrelevant = groups["other_product_line"], groups["not_relevant"]
    stats = "".join(f"<div class='stat'><b>{n}</b>{label}</div>" for n, label in
                    [(len(items), "tenders screened"), (len(power), "for SBU-2"), (len(manual), "to check"),
                     (len(other), "other teams"), (len(irrelevant), "not relevant")])
    parts = [
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width,initial-scale=1'>",
        f"<title>Tender shortlist {today:%d %b %Y}</title><style>{CSS}{EXTRA_CSS}</style></head><body>",
        "<main class='wrap'>",
        f"<section class='card'><div class='eyebrow'>BidAssist screening</div><h1>Tender shortlist, "
        f"{today:%d %b %Y}</h1><div class='summary-row'>{stats}</div></section>",
        f"<section class='card'><h2>SBU-2 power supply candidates <span class='count'>{len(power)}</span></h2>"
        + ("".join(_tender_card(s) for s in power) or "<p class='none'>None today.</p>") + "</section>",
    ]
    if manual:
        parts.append(f"<section class='card ask'><h2>Needs manual check <span class='count'>{len(manual)}</span>"
                     f"</h2>{''.join(_tender_card(s) for s in manual)}</section>")
    parts.append(f"<section class='card'><h2>Other product lines <span class='count'>{len(other)}</span></h2>"
                 + (_table(sorted(other, key=lambda s: s.route_to), True) if other else "<p class='none'>None.</p>")
                 + "</section>")
    parts.append(f"<section class='card'><details><summary>Not relevant <span class='count'>{len(irrelevant)}"
                 f"</span></summary>{_table(irrelevant, False) if irrelevant else ''}</details></section>")
    parts.append("</main></body></html>")
    return "\n".join(parts)


TRACKER_COLUMNS = ["Category", "Route to", "Confidence", "Reason", "Days left", "Tender Id", "Tender No",
                   "Authority", "Summary", "Closing Date", "Value", "EMD", "Location", "Keyword", "Corrigendum",
                   "Source", "Decision", "Notes"]


def build_tracker(items: list[Shortlisted]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Shortlist"
    ws.append(TRACKER_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for s in sorted(items, key=lambda s: (CATEGORY_ORDER.get(s.category, 9), _sort_key(s))):
        t = s.tender
        ws.append([s.category, s.route_to, s.confidence, s.reason, s.days_left, t.tender_id, t.tender_no,
                   t.authority, t.summary, t.closing_date, t.amount, t.emd, t.location, t.keyword,
                   "Y" if t.corrigendum else "N", t.source, None, None])
    decision_col = TRACKER_COLUMNS.index("Decision") + 1
    letter = ws.cell(1, decision_col).column_letter
    dv = DataValidation(type="list", formula1='"Fetch,Skip"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")
    yellow = PatternFill("solid", fgColor="FFFF00")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == "sbu2_power":
            ws.cell(r, decision_col).fill = yellow
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFGHI", (18, 18, 11, 50, 9, 22, 22, 30, 50)):
        ws.column_dimensions[col].width = width
    return wb
