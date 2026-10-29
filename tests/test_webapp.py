import io
import json
import zipfile
from datetime import date

import openpyxl
import pytest

import draft
import shortlist
import triage
import webapp
from bidtriage.analyzer import AnalysisError
from bidtriage.bidassist import Tender
from bidtriage.shortlist_out import build_tracker
from bidtriage.shortlister import Shortlisted
from bidtriage.web import jobs
from tests.fixtures import sample_assessment

RESPONSE_FILES = {"Technical_Compliance.docx": 47, "Compliance_Statement.docx": 10, "Price_Bid.xlsx": 8,
                  "Checklist.md": None}


def fake_triage(path, *, out_dir, **kwargs):
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "card.html").write_text("<html><body>CARD for HAL AHRU</body></html>", encoding="utf-8")
    (out_dir / "assessment.json").write_text(sample_assessment().model_dump_json(), encoding="utf-8")
    return out_dir


def fake_draft(path, card_dir, **kwargs):
    out = card_dir / "response"
    out.mkdir(parents=True, exist_ok=True)
    for name in RESPONSE_FILES:
        (out / name).write_bytes(b"content of " + name.encode())
    (out / "summary.json").write_text(json.dumps(RESPONSE_FILES), encoding="utf-8")
    return RESPONSE_FILES


ITEMS = [
    Shortlisted(Tender(tender_id="GEM/2026/B/8082649", authority="BEL", summary="Variable SMPS",
                       closing_date="05-Oct-2026"), "sbu2_power", "SBU-2", "high", "SMPS unit", 3),
    Shortlisted(Tender(tender_id="MES1", authority="MES", summary="Power supply to barriers"),
                "not_relevant", "none", "high", "Civil works", 2),
]


def fake_shortlist(paths, *, out_dir, **kwargs):
    out_dir.mkdir(parents=True, exist_ok=True)
    build_tracker(ITEMS).save(out_dir / "shortlist.xlsx")
    (out_dir / "digest.html").write_text("<html>digest</html>", encoding="utf-8")
    data = [{"tender_id": s.tender.tender_id, "authority": s.tender.authority, "summary": s.tender.summary,
             "closing_date": s.tender.closing_date, "category": s.category, "route_to": s.route_to,
             "confidence": s.confidence, "reason": s.reason, "days_left": s.days_left, "location": "",
             "amount": "", "emd": "", "tender_no": "", "corrigendum": False} for s in ITEMS]
    (out_dir / "shortlist.json").write_text(json.dumps(data), encoding="utf-8")
    return out_dir, ITEMS


@pytest.fixture
def env(tmp_path, monkeypatch):
    samples = tmp_path / "samples"
    (samples / "ahru-hal-hyd").mkdir(parents=True)
    (samples / "ahru-hal-hyd" / "email.txt").write_text("Subject: AHRU\n\nCheck.", encoding="utf-8")
    (samples / "ahru-hal-hyd" / "rfq.pdf").write_bytes(b"%PDF")
    monkeypatch.setattr(triage, "run_triage", fake_triage)
    monkeypatch.setattr(draft, "run_draft", fake_draft)
    monkeypatch.setattr(shortlist, "run_shortlist", fake_shortlist)
    app = webapp.create_app(data_dir=tmp_path / "data", samples_dir=samples, sync_jobs=True,
                            allowed_hosts={"localhost"})
    return app.test_client(), tmp_path / "data"


def _new_bid(client, **data):
    form = {"provider": "gemini", "subject": "HAL AHRU", "body": "Check feasibility",
            "files": [(io.BytesIO(b"%PDF-1.4"), "RFQ Tender.pdf")]} | data
    resp = client.post("/bids", data=form, content_type="multipart/form-data")
    assert resp.status_code == 302, resp.data
    return resp.headers["Location"]


def test_dashboard_renders(env):
    client, _ = env
    html = client.get("/").get_data(as_text=True)
    assert "Needs action" in html and "Screen BidAssist exports" in html
    assert "ahru-hal-hyd" in client.get("/bids/new").get_data(as_text=True)


def test_upload_bid_runs_triage_and_shows_card(env):
    client, data = env
    url = _new_bid(client)
    html = client.get(url).get_data(as_text=True)
    assert "Generate response" in html and "/card" in html
    bid_dir = next((data / "bids").iterdir())
    assert (bid_dir / "input" / "RFQ_Tender.pdf").is_file()
    assert "Subject: HAL AHRU" in (bid_dir / "input" / "email.txt").read_text(encoding="utf-8")
    assert "CARD for HAL AHRU" in client.get(url + "/card").get_data(as_text=True)


def test_bid_from_sample_folder(env):
    client, data = env
    resp = client.post("/bids", data={"provider": "gemini", "sample": "ahru-hal-hyd"},
                       content_type="multipart/form-data")
    assert resp.status_code == 302
    bid_dir = next((data / "bids").iterdir())
    assert (bid_dir / "input" / "rfq.pdf").is_file()


def test_upload_without_files_errors(env):
    client, data = env
    resp = client.post("/bids", data={"provider": "gemini", "subject": "x"}, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert "Add the bid documents" in resp.get_data(as_text=True)
    assert not (data / "bids").exists() or not any((data / "bids").iterdir())


def test_generate_response_lists_files_and_downloads(env):
    client, _ = env
    url = _new_bid(client)
    assert client.post(url + "/draft", data={"provider": "gemini"}).status_code == 302
    html = client.get(url).get_data(as_text=True)
    assert "Technical_Compliance.docx" in html and "47 placeholders" in html and "8 lines to price" in html
    assert client.get(url + "/files/Price_Bid.xlsx").data == b"content of Price_Bid.xlsx"


def test_zip_download(env):
    client, _ = env
    url = _new_bid(client)
    client.post(url + "/draft", data={"provider": "gemini"})
    resp = client.get(url + "/response.zip")
    names = zipfile.ZipFile(io.BytesIO(resp.data)).namelist()
    assert set(RESPONSE_FILES) <= set(names) and "summary.json" not in names


def test_download_rejects_traversal(env):
    client, _ = env
    url = _new_bid(client)
    client.post(url + "/draft", data={"provider": "gemini"})
    assert client.get(url + "/files/..%2F..%2Finput%2Femail.txt").status_code == 404
    assert client.get(url + "/files/../assessment.json").status_code == 404
    assert client.get("/bids/..%2F..%2Fetc/card").status_code == 404


def test_failed_job_shows_error(env, monkeypatch):
    client, _ = env

    def boom(*a, **k):
        raise AnalysisError("Gemini API error 503: high demand")

    monkeypatch.setattr(triage, "run_triage", boom)
    html = client.get(_new_bid(client)).get_data(as_text=True)
    assert "Gemini API error 503: high demand" in html
    assert 'http-equiv="refresh"' not in html


def test_running_job_page_refreshes(env):
    client, data = env
    url = _new_bid(client)
    bid_dir = next((data / "bids").iterdir())
    jobs.write_status(bid_dir, {"state": "running", "step": "draft"})
    html = client.get(url).get_data(as_text=True)
    assert 'http-equiv="refresh"' in html and "Drafting the response" in html


def test_double_submit_starts_one_job(env, monkeypatch):
    client, data = env
    url = _new_bid(client)
    bid_dir = next((data / "bids").iterdir())
    jobs.write_status(bid_dir, {"state": "running", "step": "draft"})
    calls = []
    monkeypatch.setattr(draft, "run_draft", lambda *a, **k: calls.append(1))
    client.post(url + "/draft", data={"provider": "gemini"})
    assert calls == []


def test_shortlist_upload_and_decide(env):
    client, data = env
    resp = client.post("/shortlists", data={"provider": "gemini",
                                            "files": [(io.BytesIO(b"PK"), "Power supply-Tender List.xlsx")]},
                       content_type="multipart/form-data")
    assert resp.status_code == 302
    url = resp.headers["Location"]
    html = client.get(url).get_data(as_text=True)
    assert "Variable SMPS" in html and "Fetch documents" in html and "Civil works" in html
    assert client.post(url + "/decide", data={"tender_id": "GEM/2026/B/8082649", "decision": "Fetch"}).status_code == 302
    run_dir = next((data / "shortlists").iterdir())
    assert json.loads((run_dir / "decisions.json").read_text())["GEM/2026/B/8082649"] == "Fetch"
    ws = openpyxl.load_workbook(run_dir / "result" / "shortlist.xlsx").active
    headers = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[headers.index("Tender Id")].value == "GEM/2026/B/8082649")
    assert row[headers.index("Decision")].value == "Fetch"


def test_decide_unknown_tender_400(env):
    client, _ = env
    resp = client.post("/shortlists", data={"provider": "gemini", "files": [(io.BytesIO(b"PK"), "a.xlsx")]},
                       content_type="multipart/form-data")
    url = resp.headers["Location"]
    assert client.post(url + "/decide", data={"tender_id": "NOPE", "decision": "Fetch"}).status_code == 400
    assert client.post(url + "/decide", data={"tender_id": "GEM/2026/B/8082649", "decision": "Maybe"}
                       ).status_code == 400


def test_failed_triage_can_be_rerun(env, monkeypatch):
    client, _ = env

    def boom(*a, **k):
        raise AnalysisError("could not reach the Gemini API")

    monkeypatch.setattr(triage, "run_triage", boom)
    url = _new_bid(client)
    assert "Run triage again" in client.get(url).get_data(as_text=True)
    monkeypatch.setattr(triage, "run_triage", fake_triage)
    assert client.post(url + "/triage", data={"provider": "gemini"}).status_code == 302
    html = client.get(url).get_data(as_text=True)
    assert "Generate response" in html and "could not reach" not in html

def _eml_bytes():
    from email.message import EmailMessage

    msg = EmailMessage()
    msg["Subject"] = "AHRU requirement"
    msg.set_content("Check feasibility.")
    return msg.as_bytes()


def test_email_upload_with_typed_text_and_extra_pdf(env):
    client, data = env
    _new_bid(client, files=[(io.BytesIO(_eml_bytes()), "bid.eml"), (io.BytesIO(b"%PDF"), "extra.pdf")])
    bid_dir = next((data / "bids").iterdir())
    assert sorted(p.name for p in (bid_dir / "input").iterdir()) == ["bid.eml", "extra.pdf"]
    assert json.loads((bid_dir / "meta.json").read_text())["input"] == "input"


def test_two_email_files_rejected(env):
    client, _ = env
    resp = client.post("/bids", data={"provider": "gemini", "files": [(io.BytesIO(_eml_bytes()), "a.eml"),
                                                                      (io.BytesIO(_eml_bytes()), "b.msg")]},
                       content_type="multipart/form-data")
    assert resp.status_code == 400 and "one email file" in resp.get_data(as_text=True)


def test_upload_names_keep_extension_and_do_not_collide(env):
    client, data = env
    _new_bid(client, files=[(io.BytesIO(b"%PDF"), "निविदा.pdf"), (io.BytesIO(b"%PDF"), "RFQ (1).pdf"),
                            (io.BytesIO(b"%PDF"), "RFQ_1.pdf"), (io.BytesIO(b"x"), "email.txt")])
    names = sorted(p.name for p in (next((data / "bids").iterdir()) / "input").iterdir())
    assert names == ["RFQ_1.pdf", "RFQ_1_2.pdf", "email.txt", "email_2.txt", "upload.pdf"]


def test_foreign_host_rejected(env):
    client, _ = env
    assert client.get("/", headers={"Host": "evil.example"}).status_code == 400


def test_cross_site_post_rejected(env):
    client, data = env
    resp = client.post("/bids", data={"provider": "gemini", "sample": "ahru-hal-hyd"},
                       headers={"Origin": "https://evil.example"}, content_type="multipart/form-data")
    assert resp.status_code == 403
    assert not (data / "bids").exists() or not any((data / "bids").iterdir())
    ok = client.post("/bids", data={"provider": "gemini", "sample": "ahru-hal-hyd"},
                     headers={"Origin": "http://localhost"}, content_type="multipart/form-data")
    assert ok.status_code == 302


def test_rebuild_from_saved_draft(env, monkeypatch):
    client, data = env
    url = _new_bid(client)
    client.post(url + "/draft", data={"provider": "gemini"})
    bid_dir = next((data / "bids").iterdir())
    (bid_dir / "card" / "response" / "draft.json").write_text("{}", encoding="utf-8")
    assert "Rebuild documents" in client.get(url).get_data(as_text=True)
    seen = {}
    monkeypatch.setattr(draft, "run_draft", lambda *a, **k: seen.update(k) or RESPONSE_FILES)
    assert client.post(url + "/rebuild").status_code == 302
    assert seen["from_draft"].name == "draft.json"


def test_decide_survives_corrupt_xlsx(env):
    client, data = env
    resp = client.post("/shortlists", data={"provider": "gemini", "files": [(io.BytesIO(b"PK"), "a.xlsx")]},
                       content_type="multipart/form-data")
    url = resp.headers["Location"]
    run_dir = next((data / "shortlists").iterdir())
    (run_dir / "result" / "shortlist.xlsx").write_bytes(b"not a workbook")
    r = client.post(url + "/decide", data={"tender_id": "GEM/2026/B/8082649", "decision": "Skip"})
    assert r.status_code == 302 and "msg=" in r.headers["Location"]
    assert json.loads((run_dir / "decisions.json").read_text())["GEM/2026/B/8082649"] == "Skip"

def _shortlist_with_candidate(client):
    resp = client.post("/shortlists", data={"provider": "gemini", "files": [(io.BytesIO(b"PK"), "a.xlsx")]},
                       content_type="multipart/form-data")
    return resp.headers["Location"]


def test_every_page_has_the_main_menu(env):
    client, _ = env
    sl = _shortlist_with_candidate(client)
    bid = _new_bid(client)
    for url in ("/", "/shortlists", "/bids", "/bids/new", sl, bid):
        html = client.get(url).get_data(as_text=True)
        for target in ('href="/"', 'href="/shortlists"', 'href="/bids"', 'href="/bids/new"'):
            assert target in html, (url, target)


def test_breadcrumbs_show_where_you_are(env):
    client, _ = env
    sl = _shortlist_with_candidate(client)
    html = client.get(sl).get_data(as_text=True)
    assert 'aria-label="Breadcrumb"' in html and 'href="/shortlists">Shortlists' in html


def test_fetch_then_start_bid_links_tender_and_bid_both_ways(env):
    client, data = env
    sl = _shortlist_with_candidate(client)
    client.post(sl + "/decide", data={"tender_id": "GEM/2026/B/8082649", "decision": "Fetch"})
    page = client.get(sl).get_data(as_text=True)
    assert "Start bid" in page and "tender=GEM" in page
    run_id = sl.rstrip("/").split("/")[-1]
    form = client.get(f"/bids/new?shortlist={run_id}&tender=GEM/2026/B/8082649").get_data(as_text=True)
    assert 'value="Variable SMPS' in form and 'name="tender"' in form
    url = _new_bid(client, subject="Variable SMPS", shortlist=run_id, tender="GEM/2026/B/8082649")
    bid_page = client.get(url).get_data(as_text=True)
    assert f'href="/shortlists/{run_id}' in bid_page
    assert "Open bid" in client.get(sl).get_data(as_text=True)


def test_skipped_candidate_has_no_start_bid(env):
    client, _ = env
    sl = _shortlist_with_candidate(client)
    client.post(sl + "/decide", data={"tender_id": "GEM/2026/B/8082649", "decision": "Skip"})
    assert "tender=GEM" not in client.get(sl).get_data(as_text=True)


def test_overview_lists_what_needs_action(env):
    client, _ = env
    sl = _shortlist_with_candidate(client)
    html = client.get("/").get_data(as_text=True)
    assert "Variable SMPS" in html and "Fetch documents or skip" in html
    client.post(sl + "/decide", data={"tender_id": "GEM/2026/B/8082649", "decision": "Fetch"})
    html = client.get("/").get_data(as_text=True)
    assert "Start the bid" in html
    _new_bid(client, subject="HAL AHRU card")
    html = client.get("/").get_data(as_text=True)
    assert "HAL AHRU card" in html and "Go / No-go on the card" in html


def test_bid_decision_on_card(env):
    client, data = env
    url = _new_bid(client)
    html = client.get(url).get_data(as_text=True)
    assert "Go / No-go" in html and 'class="step current"' in html
    assert client.post(url + "/decide", data={"decision": "Go"}).status_code == 302
    html = client.get(url).get_data(as_text=True)
    assert "Decision: Go" in html
    assert client.post(url + "/decide", data={"decision": "Maybe"}).status_code == 400
    bid_dir = next((data / "bids").iterdir())
    assert json.loads((bid_dir / "decision.json").read_text())["decision"] == "Go"


def test_bid_stepper_points_to_next_action(env):
    client, _ = env
    url = _new_bid(client)
    html = client.get(url).get_data(as_text=True)
    assert html.index('class="step done"') < html.index('class="step current"')
    assert "Decide on the card" in html
    client.post(url + "/decide", data={"decision": "Go"})
    client.post(url + "/draft", data={"provider": "gemini"})
    html = client.get(url).get_data(as_text=True)
    assert "Download the pack" in html

# ---------------------------------------------------------------- mail intake + replies

SETTINGS = {"screen_alerts": ["Power supply"], "mahadevan_email": "m@rugrel.com", "signature": "On behalf of M"}


def _mail(subject, attachments):
    from email.message import EmailMessage

    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = "BD <bd@rugrel.com>"
    msg["Date"] = "Wed, 07 Oct 2026 18:06:00 +0530"
    msg.set_content("Please check the feasibility.")
    for name, data in attachments:
        msg.add_attachment(data, maintype="application", subtype="octet-stream", filename=name)
    return msg.as_bytes()


@pytest.fixture
def mail_env(tmp_path, monkeypatch):
    samples = tmp_path / "samples"
    samples.mkdir()
    monkeypatch.setattr(triage, "run_triage", fake_triage)
    monkeypatch.setattr(draft, "run_draft", fake_draft)
    monkeypatch.setattr(shortlist, "run_shortlist", fake_shortlist)
    app = webapp.create_app(data_dir=tmp_path / "data", samples_dir=samples, sync_jobs=True,
                            allowed_hosts={"localhost"}, intake_settings=SETTINGS)
    return app.test_client(), tmp_path / "data"


def _intake(client, eml, entry_id="E1", name="mail.eml"):
    return client.post("/api/intake", data={"entry_id": entry_id, "file": (io.BytesIO(eml), name)},
                       content_type="multipart/form-data")


def test_intake_bid_mail_creates_bid_and_runs_triage(mail_env):
    client, data = mail_env
    r = _intake(client, _mail("AHRU Power Supply Requirement", [("RFQ.pdf", b"%PDF-1.4")]))
    assert r.status_code == 200 and r.json["kind"] == "bid"
    bid_id = r.json["bid_id"]
    meta = json.loads((data / "bids" / bid_id / "meta.json").read_text())
    assert meta["mail"]["entry_id"] == "E1" and meta["title"] == "AHRU Power Supply Requirement"
    html = client.get(f"/bids/{bid_id}").get_data(as_text=True)
    assert "From your mail" in html and "Generate response" in html


def test_repeat_entry_id_is_idempotent(mail_env):
    client, data = mail_env
    eml = _mail("AHRU", [("RFQ.pdf", b"%PDF")])
    first, second = _intake(client, eml).json, _intake(client, eml).json
    assert first["bid_id"] == second["bid_id"]
    assert len(list((data / "bids").iterdir())) == 1


def test_intake_bidassist_screens_only_power_supply(mail_env, monkeypatch):
    client, data = mail_env
    seen = {}

    def capture(paths, **kwargs):
        seen["names"] = [p.name for p in paths]
        return fake_shortlist(paths, **kwargs)

    monkeypatch.setattr(shortlist, "run_shortlist", capture)
    r = _intake(client, _mail("BidAssist alert", [("Power supply-Tender List_26102026.xlsx", b"PK"),
                                                  ("My Preference-Tender List_26102026.xlsx", b"PK")]))
    assert r.json["kind"] == "bidassist"
    assert seen["names"] == ["Power_supply-Tender_List_26102026.xlsx"]
    assert r.json["ignored"] == ["My Preference-Tender List_26102026.xlsx"]
    assert client.get(f"/shortlists/{r.json['shortlist_id']}").status_code == 200


def test_intake_bidassist_without_power_supply_is_ignored(mail_env):
    client, data = mail_env
    r = _intake(client, _mail("BidAssist alert", [("My Preference-Tender List_26102026.xlsx", b"PK")]))
    assert r.status_code == 200 and r.json["kind"] == "ignored"
    assert not (data / "shortlists").exists() or not any((data / "shortlists").iterdir())


def test_intake_rejects_non_mail(mail_env):
    client, _ = mail_env
    r = client.post("/api/intake", data={"entry_id": "E9", "file": (io.BytesIO(b"%PDF"), "x.pdf")},
                    content_type="multipart/form-data")
    assert r.status_code == 400


def _mail_bid(client):
    return _intake(client, _mail("AHRU Power Supply Requirement", [("RFQ.pdf", b"%PDF-1.4")])).json["bid_id"]


def test_go_auto_generates_pack_then_reply_due(mail_env):
    client, data = mail_env
    bid_id = _mail_bid(client)
    client.post(f"/bids/{bid_id}/decide", data={"decision": "Go"})
    replies = client.get("/api/replies").json
    assert len(replies) == 1
    r = replies[0]
    assert r["bid_id"] == bid_id and r["entry_id"] == "E1" and r["decision"] == "Go" and r["cc"] == "m@rugrel.com"
    assert "Feasible" in r["html"] and "On behalf of M" in r["html"]
    (zip_rel,) = r["attachments"]
    assert zip_rel.endswith("_response.zip") and (data / zip_rel).is_file()


def test_go_reply_waits_for_pack(mail_env, monkeypatch):
    client, _ = mail_env
    bid_id = _mail_bid(client)

    def fail(*a, **k):
        raise AnalysisError("Gemini API error 503")

    monkeypatch.setattr(draft, "run_draft", fail)
    client.post(f"/bids/{bid_id}/decide", data={"decision": "Go"})
    assert client.get("/api/replies").json == []
    assert "Gemini API error 503" in client.get(f"/bids/{bid_id}").get_data(as_text=True)


def test_no_go_reply_due_then_marked_drafted(mail_env):
    client, _ = mail_env
    bid_id = _mail_bid(client)
    client.post(f"/bids/{bid_id}/decide", data={"decision": "No-go"})
    (r,) = client.get("/api/replies").json
    assert r["decision"] == "No-go" and r["attachments"] == [] and "Not feasible" in r["html"]
    assert client.post(f"/api/replies/{bid_id}/drafted", data={"decision": "No-go"}).status_code == 200
    assert client.get("/api/replies").json == []
    assert "Reply drafted" in client.get(f"/bids/{bid_id}").get_data(as_text=True)
    assert "send the reply" in client.get("/").get_data(as_text=True)


def test_changed_decision_redue_reply(mail_env):
    client, _ = mail_env
    bid_id = _mail_bid(client)
    client.post(f"/bids/{bid_id}/decide", data={"decision": "No-go"})
    client.post(f"/api/replies/{bid_id}/drafted", data={"decision": "No-go"})
    client.post(f"/bids/{bid_id}/decide", data={"decision": "Ask"})
    (r,) = client.get("/api/replies").json
    assert r["decision"] == "Ask"


def test_manual_bid_has_no_reply(mail_env):
    client, _ = mail_env
    url = _new_bid(client)
    client.post(url + "/decide", data={"decision": "No-go"})
    assert client.get("/api/replies").json == []

def test_drafted_for_an_outdated_decision_is_rejected(mail_env):
    client, data = mail_env
    bid_id = _mail_bid(client)
    client.post(f"/bids/{bid_id}/decide", data={"decision": "No-go"})
    client.post(f"/bids/{bid_id}/decide", data={"decision": "Ask"})        # changed while the draft was made
    r = client.post(f"/api/replies/{bid_id}/drafted", data={"decision": "No-go"})
    assert r.status_code == 409
    assert not (data / "bids" / bid_id / "reply.json").exists()
    (due,) = client.get("/api/replies").json
    assert due["decision"] == "Ask"


def test_go_after_retriage_regenerates_stale_pack(mail_env, monkeypatch):
    import os
    import time

    client, data = mail_env
    bid_id = _mail_bid(client)
    calls = []

    def counting_draft(path, card_dir, **kwargs):
        calls.append(1)
        return fake_draft(path, card_dir, **kwargs)

    monkeypatch.setattr(draft, "run_draft", counting_draft)
    client.post(f"/bids/{bid_id}/decide", data={"decision": "Go"})
    assert len(calls) == 1
    summary = data / "bids" / bid_id / "card" / "response" / "summary.json"
    old = time.time() - 3600
    os.utime(summary, (old, old))                                          # pack older than the card now
    client.post(f"/bids/{bid_id}/decide", data={"decision": "Go"})
    assert len(calls) == 2