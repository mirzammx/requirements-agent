from datetime import date

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from bidtriage.docx_out import (
    add_rich,
    all_text,
    build_checklist,
    build_compliance_statement,
    build_declarations,
    build_prebid_queries,
    build_technical_compliance,
    count_placeholders,
)
from bidtriage.profile import PROFILE_KEYS
from tests.draft_fixtures import sample_draft

TODAY = date(2026, 10, 26)
PROFILE = {k: "" for k in PROFILE_KEYS} | {"company_name": "Rug-Rel Components & Systems Pvt Ltd",
                                            "address": "No. 313, Nagarbhavi, Bengaluru"}


def test_placeholders_are_highlighted_runs():
    doc = Document()
    p = doc.add_paragraph()
    add_rich(p, "GSTIN: [[GSTIN]] here")
    assert [r.text for r in p.runs] == ["GSTIN: ", "[[GSTIN]]", " here"]
    assert p.runs[1].font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert p.runs[1].bold
    assert p.runs[0].font.highlight_color is None
    assert count_placeholders("a [[x]] b [[y: z]]") == 2


def test_technical_compliance_has_section_tables_and_count():
    doc = build_technical_compliance(sample_draft(), PROFILE)
    text = all_text(doc)
    assert "Design Compliance" in text
    assert "HAL/HYD/IMM-INDG/M1/PSM-AHRU/2026/01" in text
    assert "Part-B: Project Requirements" in text
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Sl. No", "Description", "Rug-Rel Compliance"]
    assert "[[REVIEW: 100 ms" in table.rows[2].cells[2].text
    assert doc.paragraphs[0].text.startswith(f"{count_placeholders(text)} highlighted placeholder")
    assert "Rug-Rel Components & Systems Pvt Ltd" in doc.sections[0].header.paragraphs[0].text


def test_no_form_rows_returns_none():
    d = sample_draft()
    d.compliance_form_rows = []
    assert build_compliance_statement(d, PROFILE) is None


def test_compliance_statement_uses_profile_or_placeholder():
    text = all_text(build_compliance_statement(sample_draft(), PROFILE))
    assert "Rug-Rel Components & Systems Pvt Ltd" in text
    assert "[[GSTIN]]" in text
    assert "Offer validity 365 days" in text
    assert "[[MLC %]]" in text


def test_declarations_include_requested_kinds_only():
    text = all_text(build_declarations(sample_draft(), PROFILE, TODAY))
    assert "Self-Certification" in text and "Local Content" in text
    assert "land border" in text.lower()
    assert "MSE Declaration" not in text
    assert "[[Local content %]]" in text
    assert "[[Authorised signatory name]]" in text


def test_prebid_letter_numbers_queries():
    text = all_text(build_prebid_queries(sample_draft(), PROFILE, TODAY))
    assert "Rajeev Kumar Jain" in text
    assert "1. Annex-1 Sl.20 vs 8.1(d)" in text
    assert "MIL-STD-461E or 461C" in text


def test_checklist_lists_documents():
    md = build_checklist(sample_draft())
    assert "- [ ] Technical bid in sealed envelope" in md
    assert "Superscribe RFQ no. and due date" in md
    assert "Technical_Compliance.docx" in md
